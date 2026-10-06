import io
import json
from datetime import datetime

from django.contrib.admin.views.decorators import staff_member_required
from django.http import (
    HttpResponse,
    JsonResponse,
)
from django.shortcuts import render
from django.views.decorators.http import require_http_methods

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import (
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)

from documents.models import Document

from .models import IAReglaSemantica
from .services.formateador_resultados import (
    formatear_resultado_motor,
)
from .services.motor_generico import (
    ejecutar_pregunta,
)
from .services.reglas_semanticas import (
    guardar_regla_exacta,
    guardar_regla_patron,
)
from .services.sugerencias_patron import (
    sugerir_patron,
)


# ============================================================
# MOTOR CENTRAL
# ============================================================

def _ejecutar_consulta_ia(
    pregunta,
    limite=20,
):
    """
    Punto único de entrada para las vistas web/API.

    La exportación NO vuelve a llamar esta función.
    Excel/PDF reciben la respuesta ya mostrada
    al usuario.
    """

    resultado_motor = ejecutar_pregunta(
        pregunta,
        limite=limite,
    )

    respuesta_estructurada = (
        formatear_resultado_motor(
            resultado_motor
        )
    )

    diagnostico = {
        "origen":
            resultado_motor.get(
                "origen"
            ),

        "tema":
            resultado_motor.get(
                "tema"
            ),

        "modelo":
            resultado_motor.get(
                "modelo"
            ),

        "operacion":
            resultado_motor.get(
                "operacion"
            ),

        "cantidad":
            resultado_motor.get(
                "cantidad"
            ),

        "orden":
            resultado_motor.get(
                "orden"
            ),

        "filtros_orm":
            resultado_motor.get(
                "filtros_orm"
            ),

        "total":
            resultado_motor.get(
                "total"
            ),

        "limite":
            resultado_motor.get(
                "limite"
            ),
    }

    return {
        "motor":
            resultado_motor,

        "diagnostico":
            diagnostico,

        "respuesta_estructurada":
            respuesta_estructurada,

        "respuesta_texto":
            json.dumps(
                respuesta_estructurada,
                ensure_ascii=False,
                indent=2,
                default=str,
            ),
    }


# ============================================================
# PANTALLA PRINCIPAL
# ============================================================

@staff_member_required
@require_http_methods(
    [
        "GET",
        "POST",
    ]
)
def consulta_ia(request):

    contexto = {
        "pregunta": "",
        "respuesta": "",
        "respuesta_estructurada": None,
        "plan": None,
        "resultado": None,
        "error": "",
    }

    if request.method == "POST":

        pregunta = (
            request.POST.get(
                "pregunta"
            )
            or ""
        ).strip()

        contexto[
            "pregunta"
        ] = pregunta

        if not pregunta:

            contexto[
                "error"
            ] = "Escribe una pregunta."

        else:

            try:

                salida = (
                    _ejecutar_consulta_ia(
                        pregunta
                    )
                )

                contexto[
                    "plan"
                ] = salida[
                    "diagnostico"
                ]

                contexto[
                    "resultado"
                ] = salida[
                    "diagnostico"
                ]

                contexto[
                    "respuesta_estructurada"
                ] = salida[
                    "respuesta_estructurada"
                ]

                contexto[
                    "respuesta"
                ] = salida[
                    "respuesta_texto"
                ]

            except Exception as exc:

                contexto[
                    "error"
                ] = str(exc)

    return render(
        request,
        "ia_local/index.html",
        contexto,
    )


# ============================================================
# API PRINCIPAL
# ============================================================

@staff_member_required
@require_http_methods(
    [
        "POST",
    ]
)
def consulta_ia_json(request):

    pregunta = (
        request.POST.get(
            "pregunta"
        )
        or ""
    ).strip()

    if not pregunta:

        return JsonResponse(
            {
                "ok": False,
                "error":
                    "La pregunta está vacía.",
            },
            status=400,
            json_dumps_params={
                "ensure_ascii": False
            },
        )

    try:

        salida = (
            _ejecutar_consulta_ia(
                pregunta
            )
        )

        diagnostico = salida[
            "diagnostico"
        ]

        return JsonResponse(
            {
                "ok":
                    True,

                "pregunta":
                    pregunta,

                "origen":
                    diagnostico[
                        "origen"
                    ],

                "tema":
                    diagnostico[
                        "tema"
                    ],

                "modelo":
                    diagnostico[
                        "modelo"
                    ],

                "operacion":
                    diagnostico[
                        "operacion"
                    ],

                "intencion":
                    salida[
                        "motor"
                    ].get(
                        "intencion"
                    ),

                "filtros_orm":
                    diagnostico[
                        "filtros_orm"
                    ],

                "total":
                    diagnostico[
                        "total"
                    ],

                "respuesta":
                    salida[
                        "respuesta_texto"
                    ],

                "respuesta_estructurada":
                    salida[
                        "respuesta_estructurada"
                    ],
            },
            json_dumps_params={
                "ensure_ascii": False
            },
        )

    except Exception as exc:

        return JsonResponse(
            {
                "ok": False,
                "error": str(exc),
            },
            status=400,
            json_dumps_params={
                "ensure_ascii": False
            },
        )


# ============================================================
# INDEX TEST
# ============================================================

@staff_member_required
def index_test(request):

    return render(
        request,
        "ia_local/index_test.html"
    )


@staff_member_required
@require_http_methods(
    [
        "POST",
    ]
)
def consulta_test_json(request):

    pregunta = (
        request.POST.get(
            "pregunta"
        )
        or ""
    ).strip()

    if not pregunta:

        return JsonResponse(
            {
                "ok": False,
                "error":
                    "La pregunta está vacía.",
            },
            status=400,
            json_dumps_params={
                "ensure_ascii": False
            },
        )

    try:

        salida = (
            _ejecutar_consulta_ia(
                pregunta
            )
        )

        motor = salida[
            "motor"
        ]

        diagnostico = salida[
            "diagnostico"
        ]

        raw = (
            motor.get(
                "intencion_raw"
            )
            or motor.get(
                "intencion"
            )
        )

        normalizado = (
            motor.get(
                "intencion"
            )
        )

        sugerencia_patron = (
            sugerir_patron(
                pregunta,
                normalizado,
            )
        )

        return JsonResponse(
            {
                "ok":
                    True,

                "pregunta":
                    pregunta,

                "origen_interpretacion":
                    diagnostico[
                        "origen"
                    ],

                "raw":
                    raw,

                "normalizado":
                    normalizado,

                "sugerencia_patron":
                    sugerencia_patron,

                "resultado":
                    diagnostico,

                "respuesta":
                    salida[
                        "respuesta_texto"
                    ],

                "respuesta_estructurada":
                    salida[
                        "respuesta_estructurada"
                    ],
            },
            json_dumps_params={
                "ensure_ascii": False
            },
        )

    except Exception as exc:

        return JsonResponse(
            {
                "ok": False,
                "error": str(exc),
            },
            status=400,
            json_dumps_params={
                "ensure_ascii": False
            },
        )


# ============================================================
# REGLA EXACTA
# ============================================================

@staff_member_required
@require_http_methods(
    [
        "POST",
    ]
)
def guardar_regla_test_json(
    request,
):

    pregunta = (
        request.POST.get(
            "pregunta"
        )
        or ""
    ).strip()

    intencion_json = (
        request.POST.get(
            "intencion"
        )
        or ""
    ).strip()

    if not pregunta:

        return JsonResponse(
            {
                "ok": False,
                "error":
                    "La pregunta está vacía.",
            },
            status=400,
        )

    if not intencion_json:

        return JsonResponse(
            {
                "ok": False,
                "error":
                    "No existe intención "
                    "normalizada para guardar.",
            },
            status=400,
        )

    try:

        intencion = json.loads(
            intencion_json
        )

    except json.JSONDecodeError:

        return JsonResponse(
            {
                "ok": False,
                "error":
                    "La intención recibida "
                    "no es JSON válido.",
            },
            status=400,
        )

    try:

        if not isinstance(
            intencion,
            dict,
        ):

            raise ValueError(
                "La intención debe ser "
                "un objeto JSON."
            )

        if not intencion.get(
            "tema"
        ):

            raise ValueError(
                "La intención no contiene "
                "'tema'."
            )

        regla, creada = (
            guardar_regla_exacta(
                pregunta=pregunta,
                intencion=intencion,
                observacion=(
                    "Regla guardada manualmente "
                    "desde la consola de diagnóstico."
                ),
            )
        )

        return JsonResponse(
            {
                "ok":
                    True,

                "creada":
                    creada,

                "regla":
                    {
                        "id":
                            regla.pk,

                        "pregunta":
                            regla.pregunta,

                        "tema":
                            regla.tema,

                        "operacion":
                            regla.operacion,

                        "filtros":
                            regla.filtros,

                        "cantidad":
                            regla.cantidad,

                        "orden":
                            regla.orden,
                    },

                "mensaje":
                    (
                        "Regla semántica creada "
                        "correctamente."
                        if creada
                        else
                        "Regla semántica actualizada "
                        "correctamente."
                    ),
            },
            json_dumps_params={
                "ensure_ascii": False
            },
        )

    except Exception as exc:

        return JsonResponse(
            {
                "ok": False,
                "error": str(exc),
            },
            status=400,
        )


# ============================================================
# PATRONES
# ============================================================

@staff_member_required
@require_http_methods(
    [
        "POST",
    ]
)
def guardar_patron_test_json(
    request,
):

    patron = (
        request.POST.get(
            "patron"
        )
        or ""
    ).strip()

    intencion_json = (
        request.POST.get(
            "intencion"
        )
        or ""
    ).strip()

    if not patron:

        return JsonResponse(
            {
                "ok": False,
                "error":
                    "El patrón está vacío.",
            },
            status=400,
        )

    if not intencion_json:

        return JsonResponse(
            {
                "ok": False,
                "error":
                    "La intención del patrón "
                    "está vacía.",
            },
            status=400,
        )

    try:

        intencion = json.loads(
            intencion_json
        )

    except json.JSONDecodeError:

        return JsonResponse(
            {
                "ok": False,
                "error":
                    "La intención no contiene "
                    "JSON válido.",
            },
            status=400,
        )

    try:

        if not isinstance(
            intencion,
            dict,
        ):

            raise ValueError(
                "La intención debe ser "
                "un objeto JSON."
            )

        if not intencion.get(
            "tema"
        ):

            raise ValueError(
                "La intención debe contener "
                "'tema'."
            )

        regla, creada = (
            guardar_regla_patron(
                patron=patron,
                intencion=intencion,
                observacion=(
                    "Patrón guardado manualmente "
                    "desde la consola de diagnóstico."
                ),
            )
        )

        return JsonResponse(
            {
                "ok":
                    True,

                "creada":
                    creada,

                "regla":
                    {
                        "id":
                            regla.pk,

                        "tipo":
                            regla.tipo,

                        "patron":
                            regla.pregunta,

                        "patron_normalizado":
                            regla.pregunta_normalizada,

                        "tema":
                            regla.tema,

                        "operacion":
                            regla.operacion,

                        "filtros":
                            regla.filtros,

                        "cantidad":
                            regla.cantidad,

                        "orden":
                            regla.orden,
                    },

                "mensaje":
                    (
                        "Patrón semántico creado "
                        "correctamente."
                        if creada
                        else
                        "Patrón semántico actualizado "
                        "correctamente."
                    ),
            },
            json_dumps_params={
                "ensure_ascii": False
            },
        )

    except Exception as exc:

        return JsonResponse(
            {
                "ok": False,
                "error": str(exc),
            },
            status=400,
        )


# ============================================================
# DIAGNOSTICO DOCUMENTOS
# ============================================================

def diagnostico_documentos(
    request,
):

    docs = (
        Document.objects
        .select_related(
            "project",
            "company",
            "process",
            "doc_type",
            "folder",
            "sello_verificado_por",
            "liberacion_registrada_por",
        )
        .order_by(
            "-date",
            "-id",
        )[:50]
    )

    datos = []

    for doc in docs:

        datos.append(
            {
                "id":
                    doc.id,

                "number":
                    doc.number,

                "code":
                    doc.code,

                "title":
                    doc.title,

                "description":
                    doc.description,

                "revision":
                    doc.revision,

                "date":
                    (
                        str(
                            doc.date
                        )
                        if doc.date
                        else None
                    ),

                "status":
                    doc.status,

                "status_nombre":
                    doc.get_status_display(),

                "informado":
                    doc.informado,

                "informado_nombre":
                    doc.get_informado_display(),

                "project_code":
                    (
                        doc.project.code
                        if doc.project
                        else None
                    ),

                "project_name":
                    (
                        doc.project.name
                        if doc.project
                        else None
                    ),

                "company_code":
                    (
                        doc.company.code
                        if doc.company
                        else None
                    ),

                "company_name":
                    (
                        doc.company.name
                        if doc.company
                        else None
                    ),

                "process_code":
                    (
                        doc.process.code
                        if doc.process
                        else None
                    ),

                "process_name":
                    (
                        doc.process.name
                        if doc.process
                        else None
                    ),

                "doc_type_code":
                    (
                        doc.doc_type.code
                        if doc.doc_type
                        else None
                    ),

                "doc_type_name":
                    (
                        doc.doc_type.name
                        if doc.doc_type
                        else None
                    ),

                "folder_code":
                    (
                        doc.folder.code
                        if doc.folder
                        else None
                    ),

                "file":
                    (
                        doc.file.name
                        if doc.file
                        else None
                    ),

                "content_extract":
                    (
                        doc.content_extract[
                            :1000
                        ]
                        if doc.content_extract
                        else ""
                    ),

                "estado_liberacion":
                    doc.estado_liberacion,

                "estado_liberacion_nombre":
                    doc.get_estado_liberacion_display(),

                "fecha_liberacion":
                    (
                        str(
                            doc.fecha_liberacion
                        )
                        if doc.fecha_liberacion
                        else None
                    ),

                "created_at":
                    (
                        doc.created_at.isoformat()
                        if doc.created_at
                        else None
                    ),

                "updated_at":
                    (
                        doc.updated_at.isoformat()
                        if doc.updated_at
                        else None
                    ),
            }
        )

    return JsonResponse(
        {
            "total_mostrados":
                len(datos),

            "documentos":
                datos,
        },
        json_dumps_params={
            "indent": 2,
            "ensure_ascii": False,
        },
    )


# ============================================================
# REGLAS SEMANTICAS
# ============================================================

@staff_member_required
@require_http_methods(
    [
        "GET",
    ]
)
def reglas_semanticas(
    request,
):

    reglas = (
        IAReglaSemantica.objects
        .all()
        .order_by(
            "-activa",
            "tipo",
            "-veces_usada",
            "-updated_at",
        )
    )

    return render(
        request,
        "ia_local/reglas_semanticas.html",
        {
            "reglas":
                reglas,
        },
    )


@staff_member_required
@require_http_methods(
    [
        "POST",
    ]
)
def cambiar_estado_regla(
    request,
    pk,
):

    try:

        regla = (
            IAReglaSemantica.objects.get(
                pk=pk
            )
        )

        regla.activa = not regla.activa

        regla.save(
            update_fields=[
                "activa",
                "updated_at",
            ]
        )

        return JsonResponse(
            {
                "ok":
                    True,

                "id":
                    regla.pk,

                "activa":
                    regla.activa,

                "mensaje":
                    (
                        "Regla activada."
                        if regla.activa
                        else
                        "Regla desactivada."
                    ),
            }
        )

    except IAReglaSemantica.DoesNotExist:

        return JsonResponse(
            {
                "ok": False,
                "error":
                    "La regla no existe.",
            },
            status=404,
        )


# ============================================================
# EXPORTACION
# ============================================================

def _leer_respuesta_exportacion(
    request,
):
    """
    Lee la respuesta estructurada que ya está visible
    en index.html.

    No vuelve a ejecutar Qwen.
    """

    pregunta = (
        request.POST.get(
            "pregunta"
        )
        or ""
    ).strip()

    respuesta_json = (
        request.POST.get(
            "respuesta_json"
        )
        or ""
    ).strip()

    if not respuesta_json:

        raise ValueError(
            "No existe un resultado "
            "para exportar."
        )

    try:

        respuesta = json.loads(
            respuesta_json
        )

    except json.JSONDecodeError:

        raise ValueError(
            "El resultado recibido "
            "no contiene JSON válido."
        )

    if not isinstance(
        respuesta,
        dict,
    ):

        raise ValueError(
            "La respuesta debe ser "
            "un objeto JSON."
        )

    return (
        pregunta,
        respuesta,
    )


def _valor_exportable(
    valor,
):
    if valor is None:
        return ""

    if isinstance(
        valor,
        (
            dict,
            list,
            tuple,
        ),
    ):

        return json.dumps(
            valor,
            ensure_ascii=False,
            default=str,
        )

    return str(
        valor
    )


def _extraer_resultados(
    respuesta,
):
    """
    Obtiene el listado principal sin depender
    de un único tipo de consulta.
    """

    resultados = respuesta.get(
        "resultados"
    )

    if isinstance(
        resultados,
        list,
    ):
        return resultados

    return []


def _campos_resultados(
    resultados,
):
    """
    Unión de las columnas encontradas.
    Mantiene el orden de aparición.
    """

    campos = []

    for fila in resultados:

        if not isinstance(
            fila,
            dict,
        ):
            continue

        for campo in fila.keys():

            if campo not in campos:
                campos.append(
                    campo
                )

    return campos


def _nombre_archivo(
    extension,
):
    ahora = datetime.now().strftime(
        "%Y%m%d_%H%M%S"
    )

    return (
        f"consulta_ia_{ahora}."
        f"{extension}"
    )


# ============================================================
# EXCEL
# ============================================================

@staff_member_required
@require_http_methods(
    [
        "POST",
    ]
)
def exportar_resultado_excel(
    request,
):
    try:

        pregunta, respuesta = (
            _leer_respuesta_exportacion(
                request
            )
        )

    except ValueError as exc:

        return HttpResponse(
            str(exc),
            status=400,
            content_type=(
                "text/plain; "
                "charset=utf-8"
            ),
        )

    workbook = Workbook()

    # --------------------------------------------------------
    # RESUMEN
    # --------------------------------------------------------

    hoja_resumen = workbook.active
    hoja_resumen.title = "Resumen"

    hoja_resumen[
        "A1"
    ] = "Consulta IA"

    hoja_resumen[
        "A1"
    ].font = Font(
        bold=True,
        size=14,
    )

    hoja_resumen[
        "A3"
    ] = "Pregunta"

    hoja_resumen[
        "B3"
    ] = pregunta

    fila = 5

    for clave, valor in respuesta.items():

        if clave == "resultados":
            continue

        hoja_resumen.cell(
            row=fila,
            column=1,
            value=str(
                clave
            ),
        )

        hoja_resumen.cell(
            row=fila,
            column=2,
            value=_valor_exportable(
                valor
            ),
        )

        fila += 1

    hoja_resumen.column_dimensions[
        "A"
    ].width = 30

    hoja_resumen.column_dimensions[
        "B"
    ].width = 90

    for row in hoja_resumen.iter_rows():

        for cell in row:

            cell.alignment = Alignment(
                vertical="top",
                wrap_text=True,
            )

    # --------------------------------------------------------
    # RESULTADOS
    # --------------------------------------------------------

    resultados = (
        _extraer_resultados(
            respuesta
        )
    )

    if resultados:

        hoja = workbook.create_sheet(
            "Resultados"
        )

        campos = (
            _campos_resultados(
                resultados
            )
        )

        for columna, campo in enumerate(
            campos,
            start=1,
        ):

            celda = hoja.cell(
                row=1,
                column=columna,
                value=campo,
            )

            celda.font = Font(
                bold=True
            )

            celda.alignment = Alignment(
                horizontal="center",
                vertical="center",
                wrap_text=True,
            )

        for numero_fila, registro in enumerate(
            resultados,
            start=2,
        ):

            if not isinstance(
                registro,
                dict,
            ):
                registro = {
                    "resultado":
                        registro
                }

            for numero_columna, campo in enumerate(
                campos,
                start=1,
            ):

                valor = (
                    registro.get(
                        campo
                    )
                )

                celda = hoja.cell(
                    row=numero_fila,
                    column=numero_columna,
                    value=_valor_exportable(
                        valor
                    ),
                )

                celda.alignment = Alignment(
                    vertical="top",
                    wrap_text=True,
                )

        hoja.freeze_panes = "A2"

        hoja.auto_filter.ref = (
            hoja.dimensions
        )

        for numero_columna, campo in enumerate(
            campos,
            start=1,
        ):

            ancho = max(
                12,
                min(
                    45,
                    len(
                        str(
                            campo
                        )
                    ) + 5,
                ),
            )

            if campo in {
                "archivo",
                "title",
                "description",
                "texto_fila",
            }:

                ancho = 45

            hoja.column_dimensions[
                get_column_letter(
                    numero_columna
                )
            ].width = ancho

    salida = io.BytesIO()

    workbook.save(
        salida
    )

    salida.seek(
        0
    )

    response = HttpResponse(
        salida.getvalue(),
        content_type=(
            "application/"
            "vnd.openxmlformats-officedocument."
            "spreadsheetml.sheet"
        ),
    )

    response[
        "Content-Disposition"
    ] = (
        'attachment; filename="'
        + _nombre_archivo(
            "xlsx"
        )
        + '"'
    )

    return response


# ============================================================
# PDF
# ============================================================

def _columnas_pdf(
    resultados,
):
    """
    El PDF no puede mostrar 15 o 20 columnas
    cómodamente.

    Se priorizan campos funcionales.
    """

    disponibles = set(
        _campos_resultados(
            resultados
        )
    )

    prioridad = [
        "item",
        "archivo",
        "fecha_archivo",
        "estado",
        "transmittal",
        "transmittal_number",
        "codigo_documento",

        "id",
        "code",
        "number",
        "title",
        "revision",
        "date",
        "status",
        "status_nombre",

        "name",
        "description",
    ]

    columnas = [
        campo
        for campo in prioridad
        if campo in disponibles
    ]

    if columnas:

        return columnas[
            :7
        ]

    return (
        _campos_resultados(
            resultados
        )[
            :7
        ]
    )


@staff_member_required
@require_http_methods(
    [
        "POST",
    ]
)
def exportar_resultado_pdf(
    request,
):
    try:

        pregunta, respuesta = (
            _leer_respuesta_exportacion(
                request
            )
        )

    except ValueError as exc:

        return HttpResponse(
            str(exc),
            status=400,
            content_type=(
                "text/plain; "
                "charset=utf-8"
            ),
        )

    buffer = io.BytesIO()

    documento = SimpleDocTemplate(
        buffer,
        pagesize=landscape(
            A4
        ),
        rightMargin=10 * mm,
        leftMargin=10 * mm,
        topMargin=12 * mm,
        bottomMargin=12 * mm,
        title="Resultado consulta IA",
    )

    estilos = getSampleStyleSheet()

    elementos = []

    elementos.append(
        Paragraph(
            "Resultado de consulta IA",
            estilos[
                "Title"
            ],
        )
    )

    elementos.append(
        Spacer(
            1,
            5 * mm,
        )
    )

    elementos.append(
        Paragraph(
            "<b>Pregunta:</b> "
            + pregunta.replace(
                "&",
                "&amp;",
            ).replace(
                "<",
                "&lt;",
            ).replace(
                ">",
                "&gt;",
            ),
            estilos[
                "BodyText"
            ],
        )
    )

    elementos.append(
        Spacer(
            1,
            4 * mm,
        )
    )

    # --------------------------------------------------------
    # Resumen
    # --------------------------------------------------------

    resumen = []

    for clave, valor in respuesta.items():

        if clave == "resultados":
            continue

        resumen.append(
            [
                Paragraph(
                    f"<b>{clave}</b>",
                    estilos[
                        "BodyText"
                    ],
                ),
                Paragraph(
                    _valor_exportable(
                        valor
                    ),
                    estilos[
                        "BodyText"
                    ],
                ),
            ]
        )

    if resumen:

        tabla_resumen = Table(
            resumen,
            colWidths=[
                45 * mm,
                205 * mm,
            ],
            repeatRows=0,
        )

        tabla_resumen.setStyle(
            TableStyle(
                [
                    (
                        "VALIGN",
                        (
                            0,
                            0,
                        ),
                        (
                            -1,
                            -1,
                        ),
                        "TOP",
                    ),
                    (
                        "GRID",
                        (
                            0,
                            0,
                        ),
                        (
                            -1,
                            -1,
                        ),
                        0.25,
                        colors.grey,
                    ),
                    (
                        "BACKGROUND",
                        (
                            0,
                            0,
                        ),
                        (
                            0,
                            -1,
                        ),
                        colors.whitesmoke,
                    ),
                    (
                        "LEFTPADDING",
                        (
                            0,
                            0,
                        ),
                        (
                            -1,
                            -1,
                        ),
                        4,
                    ),
                    (
                        "RIGHTPADDING",
                        (
                            0,
                            0,
                        ),
                        (
                            -1,
                            -1,
                        ),
                        4,
                    ),
                ]
            )
        )

        elementos.append(
            tabla_resumen
        )

        elementos.append(
            Spacer(
                1,
                6 * mm,
            )
        )

    # --------------------------------------------------------
    # Resultados
    # --------------------------------------------------------

    resultados = (
        _extraer_resultados(
            respuesta
        )
    )

    if resultados:

        columnas = (
            _columnas_pdf(
                resultados
            )
        )

        datos = [
            [
                Paragraph(
                    f"<b>{campo}</b>",
                    estilos[
                        "BodyText"
                    ],
                )
                for campo in columnas
            ]
        ]

        for registro in resultados:

            if not isinstance(
                registro,
                dict,
            ):

                registro = {
                    "resultado":
                        registro
                }

            fila = []

            for campo in columnas:

                valor = (
                    _valor_exportable(
                        registro.get(
                            campo
                        )
                    )
                )

                valor = (
                    valor
                    .replace(
                        "&",
                        "&amp;",
                    )
                    .replace(
                        "<",
                        "&lt;",
                    )
                    .replace(
                        ">",
                        "&gt;",
                    )
                )

                fila.append(
                    Paragraph(
                        valor,
                        estilos[
                            "BodyText"
                        ],
                    )
                )

            datos.append(
                fila
            )

        ancho_total = (
            267 * mm
        )

        ancho_columna = (
            ancho_total
            / max(
                len(
                    columnas
                ),
                1,
            )
        )

        tabla = Table(
            datos,
            colWidths=[
                ancho_columna
                for _ in columnas
            ],
            repeatRows=1,
        )

        tabla.setStyle(
            TableStyle(
                [
                    (
                        "BACKGROUND",
                        (
                            0,
                            0,
                        ),
                        (
                            -1,
                            0,
                        ),
                        colors.lightgrey,
                    ),
                    (
                        "GRID",
                        (
                            0,
                            0,
                        ),
                        (
                            -1,
                            -1,
                        ),
                        0.25,
                        colors.grey,
                    ),
                    (
                        "VALIGN",
                        (
                            0,
                            0,
                        ),
                        (
                            -1,
                            -1,
                        ),
                        "TOP",
                    ),
                    (
                        "FONTSIZE",
                        (
                            0,
                            0,
                        ),
                        (
                            -1,
                            -1,
                        ),
                        7,
                    ),
                    (
                        "LEFTPADDING",
                        (
                            0,
                            0,
                        ),
                        (
                            -1,
                            -1,
                        ),
                        3,
                    ),
                    (
                        "RIGHTPADDING",
                        (
                            0,
                            0,
                        ),
                        (
                            -1,
                            -1,
                        ),
                        3,
                    ),
                ]
            )
        )

        elementos.append(
            tabla
        )

    documento.build(
        elementos
    )

    buffer.seek(
        0
    )

    response = HttpResponse(
        buffer.getvalue(),
        content_type=(
            "application/pdf"
        ),
    )

    # INLINE:
    # abre el PDF en el navegador.
    # Desde el visor puede descargarse.
    response[
        "Content-Disposition"
    ] = (
        'inline; filename="'
        + _nombre_archivo(
            "pdf"
        )
        + '"'
    )

    return response


@staff_member_required
@require_http_methods(["POST"])
def vocabulario_test_json(request):

    modelo = (
        request.POST.get("modelo")
        or ""
    ).strip()

    campo = (
        request.POST.get("campo")
        or ""
    ).strip()

    valor_canonico = (
        request.POST.get(
            "valor_canonico"
        )
        or ""
    ).strip()

    aliases_texto = (
        request.POST.get("aliases")
        or ""
    ).strip()

    modo = (
        request.POST.get("modo")
        or "probar"
    ).strip().lower()


    if not modelo:

        return JsonResponse(
            {
                "ok": False,
                "error":
                    "Debes indicar el modelo.",
            },
            status=400,
        )


    if not campo:

        return JsonResponse(
            {
                "ok": False,
                "error":
                    "Debes indicar el campo.",
            },
            status=400,
        )


    if not valor_canonico:

        return JsonResponse(
            {
                "ok": False,
                "error":
                    "Debes indicar el valor canónico.",
            },
            status=400,
        )


    aliases = [
        linea.strip()
        for linea in aliases_texto.splitlines()
        if linea.strip()
    ]


    if not aliases:

        return JsonResponse(
            {
                "ok": False,
                "error":
                    "Debes ingresar al menos un alias.",
            },
            status=400,
        )


    # ========================================================
    # PROBAR
    # ========================================================

    if modo == "probar":

        return JsonResponse(
            {
                "ok": True,

                "guardado":
                    False,

                "modelo":
                    modelo,

                "campo":
                    campo,

                "valor_canonico":
                    valor_canonico,

                "aliases":
                    aliases,

                "mensaje":
                    (
                        "Equivalencia preparada correctamente. "
                        "Todavía no se ha guardado nada."
                    ),
            },
            json_dumps_params={
                "ensure_ascii": False
            },
        )


    # ========================================================
    # GUARDAR
    # ========================================================

    if modo != "guardar":

        return JsonResponse(
            {
                "ok": False,
                "error":
                    "Modo no reconocido.",
            },
            status=400,
        )


    try:

        resultados = (
            guardar_equivalencias(
                modelo=modelo,
                campo=campo,
                valor_canonico=valor_canonico,
                aliases=aliases,
            )
        )

        return JsonResponse(
            {
                "ok": True,

                "guardado":
                    True,

                "total":
                    len(resultados),

                "equivalencias":
                    resultados,

                "mensaje":
                    (
                        f"{len(resultados)} "
                        "equivalencia(s) guardada(s) "
                        "correctamente."
                    ),
            },
            json_dumps_params={
                "ensure_ascii": False
            },
        )

    except Exception as exc:

        return JsonResponse(
            {
                "ok": False,
                "error":
                    str(exc),
            },
            status=400,
            json_dumps_params={
                "ensure_ascii": False
            },
        )